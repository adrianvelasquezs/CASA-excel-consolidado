"""
Línea de comandos. Todas las rutas se indican explícitamente.

  cronograma-neon onedrive --plantillas CARPETA_O_ARCHIVOS... --salida LIBRO.xlsx \
                           --contrasenas CONTRASENAS.csv [--profesores PROFESORES.csv] [--filas 150]
  cronograma-neon google   --plantillas CARPETA_O_ARCHIVOS... --salida PLANTILLA.xlsx \
                           [--apps-script CARPETA] [--dominio uniandes.edu.co] [--filas 150]
  cronograma-neon verificar [--informe INFORME.txt]

Sin argumentos se abre la interfaz gráfica.
"""

import argparse
import sys
import tempfile
import traceback
from pathlib import Path

from . import __version__
from .comun.plantillas import ErrorGeneracion, carpeta_plantillas_incluidas, leer_plantillas
from .google_sheets import generador as google
from .onedrive import generador as onedrive


def _onedrive(args):
    # Antes de generar, asegurar que las contraseñas se podrán guardar: un Excel sin sus claves no sirve.
    destino_claves = Path(args.contrasenas)
    try:
        destino_claves.parent.mkdir(parents=True, exist_ok=True)
        with open(destino_claves, "a", encoding="utf-8-sig"):
            pass
    except OSError as error:
        raise ErrorGeneracion(f"No se puede escribir el archivo de contraseñas {destino_claves}: {error}") from error
    existentes = onedrive.leer_contrasenas_csv(destino_claves) if destino_claves.stat().st_size else {}
    profesores = onedrive.leer_profesores_csv(args.profesores) if args.profesores else []
    claves, mensajes = onedrive.generar_excel(args.plantillas, args.salida, args.filas, profesores, existentes)
    onedrive.escribir_contrasenas_csv(args.contrasenas, claves)
    mensajes.append(f"Contraseñas guardadas en: {args.contrasenas} (no lo comparta completo)")
    return mensajes


def _google(args):
    config = {}
    if args.dominio:
        config["DOMINIO"] = args.dominio
    if args.hora is not None:
        config["HORA_RESTAURACION"] = args.hora
    if args.sin_alertas:
        config["ALERTAS_POR_CORREO"] = False
    if args.gestores_sin_filas_de_profesores:
        config["GESTOR_EDITA_FILAS_DE_PROFESORES"] = False
    mensajes = google.generar_plantilla(args.plantillas, args.salida, args.filas, config)
    if args.apps_script:
        carpeta = Path(args.apps_script)
        carpeta.mkdir(parents=True, exist_ok=True)
        google.guardar_texto(carpeta / "Codigo.gs", google.codigo_apps_script(config))
        google.guardar_texto(carpeta / "appsscript.json", google.manifiesto_apps_script())
        mensajes.append(f"Apps Script guardado en: {carpeta}")
    return mensajes


def verificar():
    """Prueba de humo del ejecutable: genera las dos versiones con las plantillas incluidas en una carpeta temporal."""
    import openpyxl

    mensajes = [f"Cronograma Neón {__version__} — verificación"]
    incluidas = carpeta_plantillas_incluidas()
    if not incluidas:
        raise ErrorGeneracion("La app no trae plantillas incluidas.")
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        primer_programa = next(iter(leer_plantillas([incluidas])[0]))
        claves, m = onedrive.generar_excel([incluidas], tmp / "onedrive.xlsx", 20, [(primer_programa, "Profesor Prueba", 3)])
        mensajes += m
        libro = openpyxl.load_workbook(tmp / "onedrive.xlsx")
        if libro.sheetnames[0] != "Instrucciones" or libro.sheetnames[-1] != "Consolidado":
            raise ErrorGeneracion(f"Hojas inesperadas en el Excel: {libro.sheetnames}")
        mensajes += google.generar_plantilla([incluidas], tmp / "google.xlsx", 20)
        if openpyxl.load_workbook(tmp / "google.xlsx").sheetnames[-1] != "Usuarios":
            raise ErrorGeneracion("La plantilla de Google no tiene la hoja Usuarios.")
        codigo = google.codigo_apps_script({"DOMINIO": "ejemplo.edu.co"})
        if "DOMINIO: 'ejemplo.edu.co'," not in codigo:
            raise ErrorGeneracion("No se aplicó la configuración del Apps Script.")
        mensajes.append(f"OK: {len(libro.sheetnames) - 2} programas, {len(claves)} contraseñas, Apps Script de {len(codigo)} caracteres.")
    return mensajes


def construir_parser():
    parser = argparse.ArgumentParser(prog="cronograma-neon", description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--version", action="version", version=__version__)
    sub = parser.add_subparsers(dest="comando", required=True)

    od = sub.add_parser("onedrive", help="Excel protegido con contraseñas para OneDrive/SharePoint")
    od.add_argument("--plantillas", nargs="+", required=True, help="carpetas y/o archivos .xlsx de las plantillas")
    od.add_argument("--salida", required=True, help="dónde guardar el Excel generado")
    od.add_argument("--contrasenas", required=True,
                    help="CSV de contraseñas: se reutiliza si existe y se actualiza con las nuevas")
    od.add_argument("--profesores", help="CSV programa,profesor,filas con filas reservadas (opcional)")
    od.add_argument("--filas", type=int, default=150, help="filas de registro por programa (por defecto 150)")

    gs = sub.add_parser("google", help="Plantilla para Google Sheets (+ Apps Script)")
    gs.add_argument("--plantillas", nargs="+", required=True, help="carpetas y/o archivos .xlsx de las plantillas")
    gs.add_argument("--salida", required=True, help="dónde guardar la plantilla generada")
    gs.add_argument("--apps-script", help="carpeta donde guardar Codigo.gs y appsscript.json (opcional)")
    gs.add_argument("--dominio", help="dominio institucional (por defecto el de Codigo.gs)")
    gs.add_argument("--hora", type=int, help="hora diaria de restauración de permisos (0-23)")
    gs.add_argument("--sin-alertas", action="store_true", help="no enviar alertas por correo a los ADMIN")
    gs.add_argument("--gestores-sin-filas-de-profesores", action="store_true",
                    help="los gestores no pueden editar las filas registradas por profesores")
    gs.add_argument("--filas", type=int, default=150, help="filas de registro por programa (por defecto 150)")

    ver = sub.add_parser("verificar", help="prueba rápida de la instalación (la usa el CI)")
    ver.add_argument("--informe", help="archivo donde escribir el resultado")
    return parser


def main(argv=None):
    for flujo in (sys.stdout, sys.stderr):  # consolas de Windows sin UTF-8: no fallar por un acento
        if hasattr(flujo, "reconfigure"):
            flujo.reconfigure(errors="replace")
    args = construir_parser().parse_args(argv)
    informe = getattr(args, "informe", None)
    try:
        mensajes = {"onedrive": _onedrive, "google": _google}.get(args.comando, lambda _a: verificar())(args)
        codigo = 0
    except ErrorGeneracion as error:
        mensajes, codigo = [f"Error: {error}"], 1
    except Exception:  # noqa: BLE001 - en la verificación se quiere el detalle completo
        mensajes, codigo = ["Error inesperado:", traceback.format_exc()], 2
    texto = "\n".join(mensajes)
    print(texto, file=sys.stdout if codigo == 0 else sys.stderr)
    if informe:
        Path(informe).write_text(texto + "\n", encoding="utf-8")
    return codigo
