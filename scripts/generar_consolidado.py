"""
Genera el Excel consolidado del Cronograma 2027 Bloque Neón a partir de las
plantillas de raw/ (que nunca se modifican).

Resultado:
  salida/Cronograma 2027 Bloque Neón - Consolidado.xlsx
  privado/contrasenas.csv   (se reutiliza si ya existe; NO compartir)

Modelo de permisos (protección de hoja + rangos con contraseña):
  - Encabezados, leyenda, columna "Programa" y hoja Consolidado: bloqueados
    (solo la contraseña de administrador desbloquea).
  - Columnas azules (Programa) de cada hoja: contraseña propia del programa.
  - Filas asignadas a un profesor en config/profesores.csv: contraseña propia
    de ese profesor (el programa ya no puede editarlas).
  - Columnas verdes (CASA) en todas las hojas: contraseña del equipo CASA.
  - Estructura del libro bloqueada (no se pueden borrar/renombrar hojas).

Uso:
  uv run scripts/generar_consolidado.py [--filas 150]
"""

import argparse
import csv
import re
import secrets
import shutil
import string
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name, xl_rowcol_to_cell

ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = ROOT / "raw"
CONFIG_DIR = ROOT / "config"
PRIVADO_DIR = ROOT / "privado"
SALIDA_DIR = ROOT / "salida"

ARCHIVO_SALIDA = SALIDA_DIR / "Cronograma 2027 Bloque Neón - Consolidado.xlsx"
ARCHIVO_PROFESORES = CONFIG_DIR / "profesores.csv"
ARCHIVO_CONTRASENAS = PRIVADO_DIR / "contrasenas.csv"

FILA_ENCABEZADO = 2  # índice 0 -> fila 3 de Excel, igual que en las plantillas
PRIMERA_FILA_DATOS = 3  # índice 0 -> fila 4 de Excel

AZUL_RAW = "FF00B0F0"
AZUL = "#00B0F0"
VERDE = "#C6E0B4"
AZUL_CLARO = "#E6F7FD"
VERDE_CLARO = "#F0F7EB"
GRIS = "#F2F2F2"

# Columnas que se llenan con fecha, listas cerradas o conteos.
COLUMNAS_FECHA = {"Fecha exacta inicio", "Fecha exacta de finalización"}
COLUMNAS_SI_NO = {
    "Montaje rúbricas Bloque Neón (SI / NO)": ["SI", "NO"],
    "Variabilidad en las notas? Responder Sí o No": ["Sí", "No"],
}
COLUMNAS_CONTEO = {
    "Número de competencias evaluadas",
    "Número de estudiantes inscritos curso",
    "Número de actividades evaluadas",
    "Número de estudiantes evaluados",
    "Número de criterios evaluados",
    "Número de objetivos evaluados",
}
# Columnas que identifican un curso: si alguna tiene datos, la fila cuenta como
# registrada (columna Programa y hoja Consolidado).
COLUMNAS_CLAVE = ("Código de curso", "CRN", "Curso")


def normalizar(texto):
    return re.sub(r"\s+", " ", str(texto)).strip()


def leer_plantillas():
    """Devuelve {programa: [(encabezado, es_azul), ...]} y el ancho de cada columna."""
    plantillas = {}
    anchos = {}
    for archivo in sorted(RAW_DIR.glob("*.xlsx")):
        if archivo.name.startswith("~$"):
            continue
        programa = archivo.name.split(" ")[0]
        ws = openpyxl.load_workbook(archivo, read_only=False).worksheets[0]
        columnas = []
        for celda in ws[FILA_ENCABEZADO + 1]:
            if celda.value is None:
                continue
            color = celda.fill.fgColor
            es_azul = bool(celda.fill.fill_type) and color.type == "rgb" and color.rgb == AZUL_RAW
            columnas.append((normalizar(celda.value), es_azul))
            ancho = ws.column_dimensions[celda.column_letter].width
            if ancho and normalizar(celda.value) not in anchos:
                anchos[normalizar(celda.value)] = ancho
        plantillas[programa] = columnas
    if not plantillas:
        sys.exit(f"No se encontraron plantillas .xlsx en {RAW_DIR}")
    return plantillas, anchos


def unificar_columnas(plantillas):
    """Une los encabezados de todas las plantillas conservando el orden original."""
    unificadas = []
    for columnas in plantillas.values():
        for i, (nombre, es_azul) in enumerate(columnas):
            existentes = [c[0] for c in unificadas]
            if nombre in existentes:
                previo = next(c for c in unificadas if c[0] == nombre)
                if previo[1] != es_azul:
                    sys.exit(f"La columna '{nombre}' tiene colores distintos entre plantillas")
                continue
            # Insertar después de la columna que la precede en esta plantilla.
            anterior = columnas[i - 1][0] if i > 0 else None
            pos = existentes.index(anterior) + 1 if anterior in existentes else len(unificadas)
            unificadas.insert(pos, (nombre, es_azul))
    return unificadas


def leer_profesores(programas, filas):
    """Lee config/profesores.csv -> {programa: [(profesor, fila_ini, fila_fin), ...]} (índices 0)."""
    asignaciones = {p: [] for p in programas}
    if not ARCHIVO_PROFESORES.exists():
        return asignaciones
    siguiente = {p: PRIMERA_FILA_DATOS for p in programas}
    with open(ARCHIVO_PROFESORES, newline="", encoding="utf-8-sig") as f:
        for n, fila in enumerate(csv.DictReader(f), start=2):
            programa = (fila.get("programa") or "").strip()
            profesor = (fila.get("profesor") or "").strip()
            cantidad = (fila.get("filas") or "").strip()
            if not programa and not profesor:
                continue
            if programa not in asignaciones:
                sys.exit(f"{ARCHIVO_PROFESORES.name} línea {n}: programa '{programa}' no existe. "
                         f"Opciones: {', '.join(programas)}")
            if not profesor or not cantidad.isdigit() or int(cantidad) < 1:
                sys.exit(f"{ARCHIVO_PROFESORES.name} línea {n}: se requiere profesor y filas (entero > 0)")
            ini = siguiente[programa]
            fin = ini + int(cantidad) - 1
            if fin >= PRIMERA_FILA_DATOS + filas:
                sys.exit(f"{ARCHIVO_PROFESORES.name}: las filas asignadas en {programa} superan "
                         f"las {filas} filas disponibles (use --filas)")
            asignaciones[programa].append((profesor, ini, fin))
            siguiente[programa] = fin + 1
    return asignaciones


def generar_contrasena():
    alfabeto = "".join(c for c in string.ascii_letters + string.digits if c not in "0O1lI")
    return "".join(secrets.choice(alfabeto) for _ in range(10))


def cargar_contrasenas(claves):
    """Reutiliza privado/contrasenas.csv y genera las que falten. claves: [(rol, ambito)]."""
    existentes = {}
    if ARCHIVO_CONTRASENAS.exists():
        with open(ARCHIVO_CONTRASENAS, newline="", encoding="utf-8-sig") as f:
            for fila in csv.DictReader(f):
                existentes[(fila["rol"], fila["ambito"])] = fila["contrasena"]
    resultado = {clave: existentes.get(clave) or generar_contrasena() for clave in claves}
    PRIVADO_DIR.mkdir(exist_ok=True)
    with open(ARCHIVO_CONTRASENAS, "w", newline="", encoding="utf-8-sig") as f:
        escritor = csv.writer(f)
        escritor.writerow(["rol", "ambito", "contrasena"])
        for (rol, ambito), contrasena in resultado.items():
            escritor.writerow([rol, ambito, contrasena])
    return resultado


def rango_columnas(indices, fila_ini, fila_fin):
    """Convierte índices de columna en un sqref con áreas contiguas: 'B4:D10 F4:K10'."""
    areas, grupo = [], []
    for i in sorted(indices):
        if grupo and i != grupo[-1] + 1:
            areas.append(grupo)
            grupo = []
        grupo.append(i)
    if grupo:
        areas.append(grupo)
    return " ".join(
        f"{xl_rowcol_to_cell(fila_ini, g[0])}:{xl_rowcol_to_cell(fila_fin, g[-1])}" for g in areas
    )


def nombre_rango(texto):
    return re.sub(r"[^A-Za-z0-9_]", "_", texto)[:200]


def referencia_hoja(hoja):
    return "'" + hoja.replace("'", "''") + "'"


def crear_formatos(wb):
    base = {"border": 1, "border_color": "#BFBFBF", "valign": "top", "text_wrap": True}
    return {
        "titulo": wb.add_format({"bold": True, "font_size": 14}),
        "nota": wb.add_format({"italic": True, "font_color": "#595959"}),
        "leyenda_azul": wb.add_format({"bg_color": AZUL, "border": 1}),
        "leyenda_verde": wb.add_format({"bg_color": VERDE, "border": 1}),
        "enc_azul": wb.add_format({**base, "bold": True, "bg_color": AZUL, "align": "center", "valign": "vcenter"}),
        "enc_verde": wb.add_format({**base, "bold": True, "bg_color": VERDE, "align": "center", "valign": "vcenter"}),
        "dato_azul": wb.add_format({**base, "bg_color": AZUL_CLARO, "locked": False}),
        "dato_verde": wb.add_format({**base, "bg_color": VERDE_CLARO, "locked": False}),
        "fecha_azul": wb.add_format({**base, "bg_color": AZUL_CLARO, "locked": False, "num_format": "dd/mm/yyyy"}),
        "fecha_verde": wb.add_format({**base, "bg_color": VERDE_CLARO, "locked": False, "num_format": "dd/mm/yyyy"}),
        "bloqueado": wb.add_format({**base, "bg_color": GRIS, "locked": True}),
        "profesor_fijo": wb.add_format({**base, "bg_color": GRIS, "locked": True, "bold": True}),
        "fecha": wb.add_format({"num_format": "dd/mm/yyyy"}),
        "h1": wb.add_format({"bold": True, "font_size": 16}),
        "h2": wb.add_format({"bold": True, "font_size": 12, "top": 1}),
        "texto": wb.add_format({"text_wrap": True, "valign": "top"}),
    }


OPCIONES_PROTECCION = {
    "select_locked_cells": True,
    "select_unlocked_cells": True,
    "autofilter": True,
    "format_columns": True,
    "format_rows": True,
}


def escribir_hoja_programa(wb, f, programa, columnas, anchos, filas, asignaciones, claves):
    ws = wb.add_worksheet(programa)
    ultima = PRIMERA_FILA_DATOS + filas - 1
    ws.write(0, 0, "Azul", f["leyenda_azul"])
    ws.write(0, 1, f"Programa / Profesores ({programa})")
    ws.write(1, 0, "Verde", f["leyenda_verde"])
    ws.write(1, 1, "Equipo CASA")
    ws.write(0, 4, f"Programa: {programa}", f["titulo"])
    ws.write(1, 4, "Gris = bloqueado. Al editar una celda de color, Excel pedirá la contraseña de su rol.", f["nota"])
    ws.set_row(0, 21.75)
    ws.set_row(1, 20.25)
    ws.set_row(FILA_ENCABEZADO, 45)

    idx = {nombre: i for i, (nombre, _) in enumerate(columnas)}
    col_programa = idx["Programa"]
    col_profesor = idx.get("Profesor")
    letras_clave = [xl_col_to_name(idx[c]) for c in COLUMNAS_CLAVE if c in idx]

    for c, (nombre, es_azul) in enumerate(columnas):
        ws.write(FILA_ENCABEZADO, c, nombre, f["enc_azul" if es_azul else "enc_verde"])
        ws.set_column(c, c, max(anchos.get(nombre, 20), 14))
        if c == col_programa:
            continue
        if nombre in COLUMNAS_FECHA:
            formato = f["fecha_azul" if es_azul else "fecha_verde"]
        else:
            formato = f["dato_azul" if es_azul else "dato_verde"]
        for r in range(PRIMERA_FILA_DATOS, ultima + 1):
            ws.write_blank(r, c, None, formato)

    # Columna Programa: fórmula bloqueada que se llena sola al registrar un curso.
    for r in range(PRIMERA_FILA_DATOS, ultima + 1):
        cond = "&".join(f"{l}{r + 1}" for l in letras_clave)
        ws.write_formula(r, col_programa, f'=IF({cond}="","","{programa}")', f["bloqueado"], "")

    # Validaciones de datos.
    for nombre, c in idx.items():
        if nombre in COLUMNAS_FECHA:
            ws.data_validation(PRIMERA_FILA_DATOS, c, ultima, c, {
                "validate": "date", "criteria": ">=", "value": "2020-01-01",
                "input_title": "Fecha", "input_message": "Formato dd/mm/aaaa",
                "error_message": "Ingrese una fecha válida (dd/mm/aaaa).",
            })
        elif nombre in COLUMNAS_SI_NO:
            ws.data_validation(PRIMERA_FILA_DATOS, c, ultima, c, {
                "validate": "list", "source": COLUMNAS_SI_NO[nombre],
            })
        elif nombre in COLUMNAS_CONTEO:
            ws.data_validation(PRIMERA_FILA_DATOS, c, ultima, c, {
                "validate": "integer", "criteria": ">=", "value": 0,
                "error_message": "Ingrese un número entero mayor o igual a 0.",
            })

    azules = [i for i, (_, es_azul) in enumerate(columnas) if es_azul and i != col_programa]
    verdes = [i for i, (_, es_azul) in enumerate(columnas) if not es_azul]

    # Filas asignadas a profesores: el nombre queda fijo y el rango tiene su propia contraseña.
    filas_asignadas = set()
    for profesor, ini, fin in asignaciones:
        filas_asignadas.update(range(ini, fin + 1))
        editables = [i for i in azules if i != col_profesor]
        if col_profesor is not None:
            for r in range(ini, fin + 1):
                ws.write(r, col_profesor, profesor, f["profesor_fijo"])
        ws.unprotect_range(
            rango_columnas(editables, ini, fin),
            nombre_rango(f"Profesor_{programa}_{profesor}"),
            claves[("profesor", f"{programa} | {profesor}")],
        )

    # Filas libres del programa: agrupar en bloques contiguos no asignados.
    libres = [r for r in range(PRIMERA_FILA_DATOS, ultima + 1) if r not in filas_asignadas]
    bloques, actual = [], []
    for r in libres:
        if actual and r != actual[-1] + 1:
            bloques.append(actual)
            actual = []
        actual.append(r)
    if actual:
        bloques.append(actual)
    if bloques:
        sqref = " ".join(rango_columnas(azules, b[0], b[-1]) for b in bloques)
        ws.unprotect_range(sqref, nombre_rango(f"Programa_{programa}"), claves[("programa", programa)])

    ws.unprotect_range(
        rango_columnas(verdes, PRIMERA_FILA_DATOS, ultima),
        nombre_rango(f"CASA_{programa}"),
        claves[("casa", "Todas las hojas")],
    )

    ws.autofilter(FILA_ENCABEZADO, 0, ultima, len(columnas) - 1)
    ws.freeze_panes(PRIMERA_FILA_DATOS, idx.get("Profesor", 4) + 1)
    ws.protect(claves[("admin", "Libro completo")], OPCIONES_PROTECCION)
    return ws


def escribir_consolidado(wb, f, programas, columnas, anchos, filas, clave_admin):
    ws = wb.add_worksheet("Consolidado")
    ultima = PRIMERA_FILA_DATOS + filas  # fila Excel (1-based) de la última fila de datos
    ultima_col = xl_col_to_name(len(columnas) - 1)
    idx = {nombre: i for i, (nombre, _) in enumerate(columnas)}

    ws.write(0, 0, "Consolidado de todos los programas (solo lectura)", f["titulo"])
    ws.write(1, 0, "Se actualiza automáticamente con lo registrado en cada hoja de programa. "
                   "Para filtrar u ordenar, copie y pegue como valores en otro libro.", f["nota"])
    ws.set_row(FILA_ENCABEZADO, 45)
    for c, (nombre, es_azul) in enumerate(columnas):
        ws.write(FILA_ENCABEZADO, c, nombre, f["enc_azul" if es_azul else "enc_verde"])
        ancho = max(anchos.get(nombre, 20), 14)
        ws.set_column(c, c, ancho, f["fecha"] if nombre in COLUMNAS_FECHA else None)

    datos, condiciones = [], []
    for p in programas:
        h = referencia_hoja(p)
        rango = f"{h}!A{PRIMERA_FILA_DATOS + 1}:{ultima_col}{ultima}"
        datos.append(f'IF({rango}="","",{rango})')
        clave = "&".join(
            f"{h}!{xl_col_to_name(idx[c])}{PRIMERA_FILA_DATOS + 1}:{xl_col_to_name(idx[c])}{ultima}"
            for c in COLUMNAS_CLAVE if c in idx
        )
        condiciones.append(f'({clave})<>""')
    formula = f'=FILTER(VSTACK({",".join(datos)}),VSTACK({",".join(condiciones)}),"Sin registros")'
    if len(formula) > 8000:
        sys.exit("La fórmula del consolidado supera el límite de Excel; reduzca --filas")
    ws.write_dynamic_array_formula(PRIMERA_FILA_DATOS, 0, PRIMERA_FILA_DATOS, 0, formula)
    ws.freeze_panes(PRIMERA_FILA_DATOS, idx.get("Profesor", 4) + 1)
    ws.protect(clave_admin, {"select_locked_cells": True, "format_columns": True})
    return ws


def escribir_instrucciones(wb, f, programas, asignaciones, clave_admin):
    ws = wb.add_worksheet("Instrucciones")
    ws.set_column(0, 0, 110)
    lineas = [
        ("h1", "Cronograma 2027 Bloque Neón — Consolidado"),
        ("texto", "Este libro reúne en un solo archivo las plantillas de todos los programas. "
                  "Cada programa tiene su propia hoja; la hoja «Consolidado» las une automáticamente."),
        ("h2", "¿Qué puede editar cada quien?"),
        ("texto", "• Celdas AZULES (columnas del programa): solo el programa dueño de la hoja, con la "
                  "contraseña de su programa. Si una fila tiene el nombre del profesor fijo (en gris), "
                  "solo ese profesor puede editarla, con su contraseña personal."),
        ("texto", "• Celdas VERDES (columnas CASA): solo el equipo CASA, con la contraseña de CASA."),
        ("texto", "• Celdas GRISES, títulos, leyendas, la columna «Programa» y la hoja «Consolidado»: "
                  "nadie las edita (solo administración)."),
        ("h2", "¿Cómo se usa?"),
        ("texto", "1. Abra su hoja (pestaña con el nombre de su programa)."),
        ("texto", "2. Escriba en una celda de su color. Excel pedirá una contraseña: ingrese la de su rol. "
                  "Queda desbloqueado mientras tenga el archivo abierto."),
        ("texto", "3. La columna «Programa» se llena sola cuando registra el código, CRN o nombre del curso."),
        ("texto", "4. Las fechas van en formato dd/mm/aaaa; las columnas SI/NO tienen lista desplegable."),
        ("h2", "Programas incluidos"),
        ("texto", ", ".join(programas)),
    ]
    con_profesores = {p: a for p, a in asignaciones.items() if a}
    if con_profesores:
        lineas.append(("h2", "Filas asignadas a profesores"))
        for p, lista in con_profesores.items():
            detalle = "; ".join(f"{prof}: filas {ini + 1}–{fin + 1}" for prof, ini, fin in lista)
            lineas.append(("texto", f"{p} → {detalle}"))
    for r, (estilo, texto) in enumerate(lineas):
        ws.write(r, 0, texto, f[estilo])
    ws.protect(clave_admin, {"select_locked_cells": True})
    return ws


def proteger_estructura(ruta, contrasena, codificar):
    """XlsxWriter no protege la estructura del libro: se inserta <workbookProtection> en workbook.xml."""
    etiqueta = f'<workbookProtection workbookPassword="{codificar(contrasena)}" lockStructure="1"/>'
    tmp = Path(tempfile.mkstemp(suffix=".xlsx")[1])
    with zipfile.ZipFile(ruta) as origen, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as destino:
        for item in origen.infolist():
            contenido = origen.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml = contenido.decode("utf-8")
                if "<bookViews>" not in xml:
                    sys.exit("No se pudo proteger la estructura del libro (workbook.xml inesperado)")
                contenido = xml.replace("<bookViews>", etiqueta + "<bookViews>", 1).encode("utf-8")
            destino.writestr(item, contenido)
    shutil.move(tmp, ruta)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--filas", type=int, default=150, help="filas de registro por programa (default 150)")
    args = parser.parse_args()

    plantillas, anchos = leer_plantillas()
    programas = list(plantillas)
    columnas = unificar_columnas(plantillas)
    for programa, cols in plantillas.items():
        faltan = [c for c, _ in columnas if c not in {n for n, _ in cols}]
        if faltan:
            print(f"Aviso: {programa} no tenía {faltan}; se agregan para unificar.")

    asignaciones = leer_profesores(programas, args.filas)
    claves_necesarias = [("admin", "Libro completo"), ("casa", "Todas las hojas")]
    claves_necesarias += [("programa", p) for p in programas]
    claves_necesarias += [("profesor", f"{p} | {prof}") for p in programas for prof, _, _ in asignaciones[p]]
    claves = cargar_contrasenas(claves_necesarias)

    SALIDA_DIR.mkdir(exist_ok=True)
    wb = xlsxwriter.Workbook(str(ARCHIVO_SALIDA), {"use_future_functions": True})
    f = crear_formatos(wb)
    escribir_instrucciones(wb, f, programas, asignaciones, claves[("admin", "Libro completo")])
    for p in programas:
        escribir_hoja_programa(wb, f, p, columnas, anchos, args.filas, asignaciones[p], claves)
    consolidado = escribir_consolidado(wb, f, programas, columnas, anchos, args.filas, claves[("admin", "Libro completo")])
    wb.close()
    proteger_estructura(ARCHIVO_SALIDA, claves[("admin", "Libro completo")], consolidado._encode_password)

    print(f"Generado: {ARCHIVO_SALIDA.relative_to(ROOT)}")
    print(f"Contraseñas: {ARCHIVO_CONTRASENAS.relative_to(ROOT)} (no compartir este archivo)")
    print(f"{len(programas)} programas, {len(columnas)} columnas, {args.filas} filas por programa, "
          f"{sum(len(a) for a in asignaciones.values())} profesores con filas asignadas")


if __name__ == "__main__":
    main()
