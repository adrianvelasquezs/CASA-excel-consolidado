"""Corre la simulación de Google Sheets (Node) contra la plantilla generada y el Codigo.gs incluido."""

import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from cronograma_neon.google_sheets import generador as google

SIMULACION = Path(__file__).resolve().parent / "apps_script" / "simulacion.js"


@pytest.mark.skipif(shutil.which("node") is None, reason="requiere Node.js")
def test_simulacion_apps_script(plantillas, tmp_path):
    destino = tmp_path / "plantilla.xlsx"
    google.generar_plantilla([plantillas], destino)
    estructura = {}
    for ws in openpyxl.load_workbook(destino).worksheets:
        estructura[ws.title] = {
            "header": [(c.value, c.fill.fgColor.rgb if c.fill.fill_type else None) for c in ws[3]],
            "max_row": ws.max_row,
            "usuarios_hdr": [c.value for c in ws[1]][:4],
        }
    archivo = tmp_path / "estructura.json"
    archivo.write_text(json.dumps(estructura, ensure_ascii=False), encoding="utf-8")

    resultado = subprocess.run(["node", str(SIMULACION), str(google.ARCHIVO_CODIGO), str(archivo)],
                               capture_output=True, text=True, encoding="utf-8")
    assert resultado.returncode == 0, resultado.stdout + resultado.stderr
    assert "Todas las pruebas pasaron" in resultado.stdout
