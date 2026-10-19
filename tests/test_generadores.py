import shutil

import openpyxl
import pytest

from cronograma_neon import cli
from cronograma_neon.comun.plantillas import ErrorGeneracion, leer_plantillas
from cronograma_neon.google_sheets import generador as google
from cronograma_neon.onedrive import generador as onedrive


def test_lee_plantillas_desde_carpeta_o_archivos(plantillas):
    desde_carpeta, _ = leer_plantillas([plantillas])
    assert len(desde_carpeta) == 12
    dos = sorted(plantillas.glob("*.xlsx"))[:2]
    desde_archivos, _ = leer_plantillas(dos)
    assert list(desde_archivos) == [p.name.split(" ")[0] for p in dos]


def test_rechaza_programa_repetido(plantillas, tmp_path):
    original = next(plantillas.glob("AFIN*.xlsx"))
    copia = tmp_path / "AFIN copia.xlsx"
    shutil.copy(original, copia)
    with pytest.raises(ErrorGeneracion, match="dos plantillas"):
        leer_plantillas([original, copia])


def test_rechaza_ruta_inexistente(tmp_path):
    with pytest.raises(ErrorGeneracion, match="No existe"):
        leer_plantillas([tmp_path / "no-existe"])


def test_excel_onedrive_reutiliza_contrasenas(plantillas, tmp_path):
    destino = tmp_path / "salida" / "libro.xlsx"  # la carpeta de destino se crea si no existe
    claves, mensajes = onedrive.generar_excel([plantillas], destino, 30, [("AFIN", "Ana", 5)])
    assert destino.exists()
    assert len(claves) == 15 and ("profesor", "AFIN | Ana") in claves
    libro = openpyxl.load_workbook(destino)
    assert libro.sheetnames[0] == "Instrucciones" and libro.sheetnames[-1] == "Consolidado"
    assert libro["AFIN"]["E4"].value == "Ana"

    otra_vez, mensajes = onedrive.generar_excel([plantillas], destino, 30, [("AFIN", "Ana", 5)], claves)
    assert otra_vez == claves
    assert any("reutilizaron" in m for m in mensajes)


def test_profesores_invalidos(plantillas, tmp_path):
    with pytest.raises(ErrorGeneracion, match="no está entre"):
        onedrive.generar_excel([plantillas], tmp_path / "x.xlsx", 30, [("NOEXISTE", "Ana", 5)])
    with pytest.raises(ErrorGeneracion, match="superan"):
        onedrive.generar_excel([plantillas], tmp_path / "x.xlsx", 30, [("AFIN", "Ana", 31)])
    assert not (tmp_path / "x.xlsx").exists()


def test_csv_ida_y_vuelta(tmp_path):
    profesores = [("AFIN", "Ana Pérez", 5), ("MF", "Luis", 2)]
    onedrive.escribir_profesores_csv(tmp_path / "p.csv", profesores)
    assert onedrive.leer_profesores_csv(tmp_path / "p.csv") == profesores
    claves = {("admin", "Libro completo"): "abc", ("programa", "AFIN"): "xyz"}
    onedrive.escribir_contrasenas_csv(tmp_path / "c.csv", claves)
    assert onedrive.leer_contrasenas_csv(tmp_path / "c.csv") == claves
    (tmp_path / "malo.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ErrorGeneracion):
        onedrive.leer_contrasenas_csv(tmp_path / "malo.csv")


def test_plantilla_google_usa_el_dominio(plantillas, tmp_path):
    destino = tmp_path / "plantilla.xlsx"
    google.generar_plantilla([plantillas], destino, 20, {"DOMINIO": "ejemplo.edu.co"})
    libro = openpyxl.load_workbook(destino)
    assert libro.sheetnames[-1] == "Usuarios"
    assert "@ejemplo.edu.co" in libro["Usuarios"]["F2"].value
    assert [c.value for c in libro["AFIN"][3]][-2:] == list(google.COLUMNAS_SISTEMA)


def test_config_apps_script():
    por_defecto = google.config_apps_script_por_defecto()
    assert google.codigo_apps_script() == google.ARCHIVO_CODIGO.read_text(encoding="utf-8")
    codigo = google.codigo_apps_script({"DOMINIO": "@Otro.edu.co", "HORA_RESTAURACION": 5, "ALERTAS_POR_CORREO": False})
    assert "DOMINIO: 'otro.edu.co'," in codigo and "HORA_RESTAURACION: 5," in codigo and "ALERTAS_POR_CORREO: false," in codigo
    assert google.config_apps_script_por_defecto() == por_defecto  # el archivo incluido no cambia
    with pytest.raises(ErrorGeneracion, match="Dominio"):
        google.codigo_apps_script({"DOMINIO": "no valido"})
    with pytest.raises(ErrorGeneracion, match="hora"):
        google.codigo_apps_script({"HORA_RESTAURACION": 24})


def test_cli_onedrive_y_google(plantillas, tmp_path):
    contrasenas = tmp_path / "privado" / "c.csv"
    contrasenas.parent.mkdir()
    assert cli.main(["onedrive", "--plantillas", str(plantillas), "--salida", str(tmp_path / "od.xlsx"),
                     "--contrasenas", str(contrasenas), "--filas", "20"]) == 0
    primera = contrasenas.read_text(encoding="utf-8-sig")
    assert cli.main(["onedrive", "--plantillas", str(plantillas), "--salida", str(tmp_path / "od.xlsx"),
                     "--contrasenas", str(contrasenas), "--filas", "20"]) == 0
    assert contrasenas.read_text(encoding="utf-8-sig") == primera  # se reutilizan

    assert cli.main(["google", "--plantillas", str(plantillas), "--salida", str(tmp_path / "gs.xlsx"),
                     "--apps-script", str(tmp_path / "as"), "--dominio", "ejemplo.edu.co"]) == 0
    assert "DOMINIO: 'ejemplo.edu.co'," in (tmp_path / "as" / "Codigo.gs").read_text(encoding="utf-8")
    assert (tmp_path / "as" / "appsscript.json").exists()


def test_cli_error_legible(tmp_path, capsys):
    assert cli.main(["google", "--plantillas", str(tmp_path / "nada"), "--salida", str(tmp_path / "x.xlsx")]) == 1
    assert "Error: No existe" in capsys.readouterr().err


def test_verificar(tmp_path):
    informe = tmp_path / "informe.txt"
    assert cli.main(["verificar", "--informe", str(informe)]) == 0
    assert "OK:" in informe.read_text(encoding="utf-8")


def test_cli_crea_carpeta_de_contrasenas(plantillas, tmp_path):
    contrasenas = tmp_path / "no" / "existe" / "c.csv"
    assert cli.main(["onedrive", "--plantillas", str(plantillas), "--salida", str(tmp_path / "od.xlsx"),
                     "--contrasenas", str(contrasenas), "--filas", "20"]) == 0
    assert len(onedrive.leer_contrasenas_csv(contrasenas)) == 14
